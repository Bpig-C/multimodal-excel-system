"""KF Excel 解析器 - 使用 pandas + openpyxl"""
from pathlib import Path
from .node import Node
import pandas as pd
import os
import json
from .excel_tools import (
    extract_excel_images, detect_excel_type,
    table_str_wps_image_name_2_markdown_img_name,
    extract_image_paths, update_column_names
)


def get_microsoft_excel_image_positions(xlsx_path):
    """获取Microsoft Excel中图片所在的行号"""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(xlsx_path)
        ws = wb.active

        image_positions = {}

        if hasattr(ws, '_images'):
            for img in ws._images:
                if hasattr(img, 'anchor') and hasattr(img.anchor, '_from'):
                    row = img.anchor._from.row

                    if hasattr(img, 'path'):
                        img_filename = os.path.basename(img.path)
                    else:
                        img_filename = f"image{row}.png"

                    if row not in image_positions:
                        image_positions[row] = []
                    image_positions[row].append(img_filename)

        wb.close()
        return image_positions

    except ImportError:
        print("警告: 未安装openpyxl，无法获取图片位置信息")
        return {}
    except Exception as e:
        print(f"获取图片位置失败: {e}")
        return {}


def insert_image_references_to_dataframe(df, image_mapping, image_column='问题图片'):
    """将图片引用插入到DataFrame的指定列"""
    if image_column not in df.columns:
        print(f"警告: 列 '{image_column}' 不存在")
        return df

    sorted_images = sorted(
        image_mapping.items(),
        key=lambda x: int(x[0].replace('image', '').replace('.png', '').replace('.jpeg', '').replace('.jpg', ''))
    )

    for idx, (original_name, uuid_name) in enumerate(sorted_images):
        if idx < len(df):
            df.at[idx, image_column] = f'=DISPIMG("{uuid_name}",1)'

    return df


def excel_to_html_table(excel_path: str, sheet_name=None, image_mapping=None) -> str:
    """使用 pandas 将 Excel 转换为 HTML 表格"""
    try:
        df = pd.read_excel(excel_path, sheet_name=sheet_name)

        if image_mapping and isinstance(image_mapping, dict) and len(image_mapping) > 0:
            df = insert_image_references_to_dataframe(df, image_mapping)

        html = df.to_html(index=False, na_rep='', border=1)
        return html
    except Exception as e:
        print(f"转换 Excel 到 HTML 失败: {e}")
        return "<table></table>"


def process_excel_file(src_file_path: str, output_dir: str = None):
    """
    处理Excel文件，提取表格和图片信息并保存为JSON格式
    """
    from .config import COLUMN_MAPPINGS

    # KF 必要列（排除 DEFAULT_VALUES 中允许缺失的可选字段）
    REQUIRED_COLUMNS = [
        '快反编号', '发生时间', '问题原因及分析', '问题图片',
        '短期改善措施', '长期改善措施', '所属分类',
        '客户名称', '产品型号', '缺陷类型不良现象',
    ]

    src_file = Path(src_file_path)
    filename = src_file.stem

    if output_dir is None:
        out_root_dir = './Datas/tree_from_excel/' + filename
    else:
        out_root_dir = output_dir

    os.makedirs(out_root_dir, exist_ok=True)

    # 提取图片
    out_img_dir = out_root_dir + '/imgs'
    image_result = extract_excel_images(src_file, out_img_dir)

    image_mapping = None
    image_files_map = {}

    if isinstance(image_result, dict):
        image_mapping = image_result
        print(f"Microsoft Excel: 提取了 {len(image_mapping)} 张图片")
    else:
        print(f"WPS Excel: 提取了 {image_result} 张图片")
        if image_result > 0:
            for fname in os.listdir(out_img_dir):
                if fname.startswith('ID_'):
                    uuid = os.path.splitext(fname)[0]
                    image_files_map[uuid] = fname

    # 读取所有工作表
    try:
        with pd.ExcelFile(src_file_path) as excel_file:
            sheet_names = excel_file.sheet_names

        # 校验必要列（先读第一个工作表做列名检查，支持别名重映射）
        df_check = pd.read_excel(src_file_path, sheet_name=sheet_names[0], nrows=0)
        actual_columns = {COLUMN_MAPPINGS.get(c, c) for c in df_check.columns}
        missing = [col for col in REQUIRED_COLUMNS if col not in actual_columns]
        if missing:
            raise ValueError(f"KF Excel 缺少必要列: {missing}，请确认上传了正确的快反记录文件。")

        node_list = []

        for sheet_name in sheet_names:
            html_table = excel_to_html_table(src_file_path, sheet_name, image_mapping)
            table_str = table_str_wps_image_name_2_markdown_img_name(html_table, image_files_map)
            imgs = extract_image_paths(table_str)

            node = Node()
            node.add_table(table_str)
            node.imgs = imgs
            node_list.append(node)

        if not node_list:
            html_table = excel_to_html_table(src_file_path, None, image_mapping)
            table_str = table_str_wps_image_name_2_markdown_img_name(html_table, image_files_map)
            imgs = extract_image_paths(table_str)

            node = Node()
            node.add_table(table_str)
            node.imgs = imgs
            node_list.append(node)

    except ValueError:
        raise
    except Exception as e:
        print(f"处理 Excel 文件失败: {e}")
        node = Node()
        node.add_table("<table></table>")
        node.imgs = []
        node_list = [node]

    # 保存为 JSON
    for i, node in enumerate(node_list):
        json_path = out_root_dir + '/' + f'page_{i}.json'
        os.makedirs(os.path.dirname(json_path), exist_ok=True)
        with open(json_path, "w", encoding="utf-8") as f:
            f.write(json.dumps(node.to_dict(), ensure_ascii=False, indent=4))
        update_column_names(json_path)

    return {
        "table_count": len(node_list),
        "output_directory": out_root_dir,
        "image_directory": out_img_dir,
        "status": "success"
    }
